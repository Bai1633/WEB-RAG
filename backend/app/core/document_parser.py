"""Document parser - multi-format document parsing with encoding detection.

Supports: PDF, DOCX, Markdown, HTML, TXT, CSV
Features: encoding detection (chardet), table preservation for PDF/DOCX
"""

from __future__ import annotations

import csv
import io
import logging
from pathlib import Path

import chardet
import markdown
from bs4 import BeautifulSoup

logger = logging.getLogger(__name__)


class ParsedDocument:
    """Result of document parsing."""

    def __init__(
        self,
        text: str,
        metadata: dict | None = None,
    ) -> None:
        self.text = text
        self.metadata = metadata or {}


class DocumentParser:
    """Parse various document formats into plain text."""

    SUPPORTED_MIME_TYPES = {
        "application/pdf": "parse_pdf",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "parse_docx",
        "text/markdown": "parse_markdown",
        "text/html": "parse_html",
        "text/plain": "parse_txt",
        "text/csv": "parse_csv",
    }

    def __init__(self) -> None:
        pass

    def supports(self, mime_type: str) -> bool:
        """Check if the MIME type is supported."""
        return mime_type in self.SUPPORTED_MIME_TYPES

    def parse(self, file_path: str, mime_type: str) -> ParsedDocument:
        """Parse a document file based on its MIME type.

        Args:
            file_path: Path to the file.
            mime_type: MIME type of the file.

        Returns:
            ParsedDocument with text and metadata.

        Raises:
            ValueError: If MIME type is not supported.
        """
        if not self.supports(mime_type):
            raise ValueError(f"Unsupported MIME type: {mime_type}")

        method_name = self.SUPPORTED_MIME_TYPES[mime_type]
        method = getattr(self, method_name)
        return method(file_path)

    def parse_txt(self, file_path: str) -> ParsedDocument:
        """Parse plain text file with encoding detection."""
        raw = Path(file_path).read_bytes()
        encoding = self._detect_encoding(raw)
        text = raw.decode(encoding, errors="replace")
        metadata = {"encoding": encoding, "parser": "txt"}
        return ParsedDocument(text=text, metadata=metadata)

    def parse_pdf(self, file_path: str) -> ParsedDocument:
        """Parse PDF file, preserving table structure where possible."""
        import pdfplumber

        text_parts: list[str] = []
        metadata: dict[str, object] = {"parser": "pdfplumber"}

        with pdfplumber.open(file_path) as pdf:
            metadata["pages"] = len(pdf.pages)

            for page_num, page in enumerate(pdf.pages, 1):
                tables = page.extract_tables()
                if tables:
                    for table in tables:
                        table_text = self._table_to_markdown(table)
                        text_parts.append(f"\n## Table (page {page_num})\n{table_text}\n")

                page_text = page.extract_text() or ""
                if page_text.strip():
                    text_parts.append(f"\n--- Page {page_num} ---\n{page_text}")

        text = "\n".join(text_parts)
        return ParsedDocument(text=text, metadata=metadata)

    def parse_docx(self, file_path: str) -> ParsedDocument:
        """Parse Word document, preserving tables."""
        from docx import Document

        doc = Document(file_path)
        text_parts: list[str] = []

        for element in doc.element.body:
            if element.tag.endswith("}p"):
                para_text = "".join(
                    node.text for node in element.iter() if node.tag.endswith("}t") and node.text
                )
                if para_text.strip():
                    text_parts.append(para_text)
            elif element.tag.endswith("}tbl"):
                table_data = []
                for row in element.iter():
                    if row.tag.endswith("}tr"):
                        row_data = []
                        for cell in row.iter():
                            if cell.tag.endswith("}tc"):
                                cell_text = "".join(
                                    t.text for t in cell.iter() if t.tag.endswith("}t") and t.text
                                )
                                row_data.append(cell_text)
                        if row_data:
                            table_data.append(row_data)
                if table_data:
                    table_md = self._table_to_markdown(table_data)
                    text_parts.append(f"\n## Table\n{table_md}\n")

        text = "\n\n".join(text_parts)
        metadata = {"parser": "python-docx", "paragraphs": len(doc.paragraphs)}
        return ParsedDocument(text=text, metadata=metadata)

    def parse_markdown(self, file_path: str) -> ParsedDocument:
        """Parse Markdown file, extract text from HTML conversion."""
        raw = Path(file_path).read_bytes()
        encoding = self._detect_encoding(raw)
        md_content = raw.decode(encoding, errors="replace")

        html = markdown.markdown(md_content, extensions=["tables", "fenced_code"])
        soup = BeautifulSoup(html, "html.parser")
        text = soup.get_text(separator="\n", strip=True)

        metadata = {"encoding": encoding, "parser": "markdown"}
        return ParsedDocument(text=text, metadata=metadata)

    def parse_html(self, file_path: str) -> ParsedDocument:
        """Parse HTML file, extract clean text."""
        raw = Path(file_path).read_bytes()
        encoding = self._detect_encoding(raw)
        html_content = raw.decode(encoding, errors="replace")

        soup = BeautifulSoup(html_content, "html.parser")

        for element in soup(["script", "style", "nav", "footer", "header"]):
            element.decompose()

        text = soup.get_text(separator="\n", strip=True)
        title = soup.title.string if soup.title else ""

        metadata = {"encoding": encoding, "title": title, "parser": "beautifulsoup"}
        return ParsedDocument(text=text, metadata=metadata)

    def parse_csv(self, file_path: str) -> ParsedDocument:
        """Parse CSV file, convert to markdown table format."""
        raw = Path(file_path).read_bytes()
        encoding = self._detect_encoding(raw)
        content = raw.decode(encoding, errors="replace")

        reader = csv.reader(io.StringIO(content))
        rows = list(reader)

        if not rows:
            return ParsedDocument(text="", metadata={"encoding": encoding, "parser": "csv"})

        table_text = self._table_to_markdown(rows)
        text = f"# CSV Data\n\n{table_text}"
        metadata = {"encoding": encoding, "parser": "csv", "rows": len(rows)}
        return ParsedDocument(text=text, metadata=metadata)

    @staticmethod
    def _detect_encoding(raw_bytes: bytes) -> str:
        """Detect text encoding using chardet."""
        result = chardet.detect(raw_bytes[:100000])
        encoding = result.get("encoding", "utf-8")
        confidence = result.get("confidence", 0)

        if confidence < 0.7 or not encoding:
            encoding = "utf-8"

        return encoding

    @staticmethod
    def _table_to_markdown(table: list[list[str]]) -> str:
        """Convert a 2D table list to Markdown table format."""
        if not table:
            return ""

        cleaned = [[str(cell).strip() if cell else "" for cell in row] for row in table]

        lines = []
        lines.append("| " + " | ".join(cleaned[0]) + " |")
        lines.append("| " + " | ".join("---" for _ in cleaned[0]) + " |")
        for row in cleaned[1:]:
            lines.append("| " + " | ".join(row) + " |")

        return "\n".join(lines)
